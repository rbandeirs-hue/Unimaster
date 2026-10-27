# -*- coding: utf-8 -*-
"""Planilha para o contador (utils/exportacao_contabil.py). Só leitura.

Uso:  .venv/bin/python -m pytest tests/test_exportacao_contabil.py -q
"""
import io
import os
import sys
from datetime import date

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402

from config import get_db_connection  # noqa: E402
from utils import despesas_pagamento as dp  # noqa: E402
from utils import exportacao_contabil as ex  # noqa: E402


def test_totais_batem_com_entradas_x_saidas():
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        for academia in (2, 14):
            hoje = date.today()
            ini, fim = date(hoje.year, hoje.month, 1), dp._fim_do_mes(hoje.year, hoje.month)
            dados, n = ex.planilha(cur, academia, "Teste", ini, fim)
            wb = load_workbook(io.BytesIO(dados))
            assert wb.sheetnames == ["Resumo", "Entradas", "Saídas", "Razão"]
            resumo = {wb["Resumo"].cell(i, 1).value: wb["Resumo"].cell(i, 2).value for i in range(5, 8)}
            es = dp.entradas_saidas(cur, academia, hoje.year, hoje.month)
            assert round(resumo["Entradas"], 2) == round(float(es["entradas"]), 2)
            assert round(resumo["Saídas"], 2) == round(float(es["saidas"]), 2)
            assert wb["Entradas"].max_row - 1 - (1 if n["entradas"] else 0) == n["entradas"]
    finally:
        conn.close()
