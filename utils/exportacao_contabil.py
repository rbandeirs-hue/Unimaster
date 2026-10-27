"""Planilha para o contador: o movimento do período em Excel.

Quatro abas, com valores numéricos (formato de moeda) para o contador somar
e conferir sem redigitar:

  Resumo    período, totais de entradas e saídas, resultado, quebra por
            categoria e por forma de pagamento
  Entradas  receitas não canceladas, pela data — com aluno, forma e conta
  Saídas    despesas pagas, pela data do pagamento — com forma e conta
  Razão     todos os lançamentos efetivados do caixa no período, por conta,
            incluindo estornos e transferências (o "extrato" consolidado)

Mesmos critérios do Entradas x Saídas (utils/despesas_pagamento.py).
"""
import io
from decimal import Decimal

FMT_MOEDA = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
FMT_DATA = "DD/MM/YYYY"


def _linhas_entradas(cur, academia_id, inicio, fim):
    cur.execute(
        """SELECT r.data, r.id, r.descricao, COALESCE(NULLIF(TRIM(r.categoria), ''), 'Sem categoria') AS categoria,
                  COALESCE(a.nome, a2.nome, a3.nome) AS aluno, fp.nome AS forma, c.nome AS conta, r.valor
           FROM receitas r
           LEFT JOIN alunos a ON a.id = r.id_aluno
           LEFT JOIN mensalidade_aluno ma ON ma.id = r.id_mensalidade_aluno
           LEFT JOIN alunos a2 ON a2.id = ma.aluno_id
           LEFT JOIN cobranca_avulsa ca ON ca.id = r.id_cobranca_avulsa
           LEFT JOIN alunos a3 ON a3.id = ca.aluno_id
           LEFT JOIN formas_pagamento fp ON fp.id = r.id_forma_pagamento
           LEFT JOIN fin_lancamentos l ON l.id = r.id_lancamento
           LEFT JOIN fin_contas c ON c.id = l.id_conta
           WHERE r.id_academia = %s AND COALESCE(r.cancelada, 0) = 0
             AND r.data BETWEEN %s AND %s
           ORDER BY r.data, r.id""", (academia_id, inicio, fim))
    return cur.fetchall()


def _linhas_saidas(cur, academia_id, inicio, fim):
    cur.execute(
        """SELECT l.data_efetivacao AS data, d.id, d.data AS vencimento, d.descricao,
                  COALESCE(NULLIF(TRIM(d.categoria), ''), 'Sem categoria') AS categoria,
                  fp.nome AS forma, c.nome AS conta, l.valor
           FROM despesas d
           JOIN fin_lancamentos l ON l.id = d.id_lancamento
                AND l.status = 'efetivado' AND l.id_estorno IS NULL
           LEFT JOIN formas_pagamento fp ON fp.id = l.id_forma_pagamento
           LEFT JOIN fin_contas c ON c.id = l.id_conta
           WHERE d.id_academia = %s AND l.data_efetivacao BETWEEN %s AND %s
           ORDER BY l.data_efetivacao, d.id""", (academia_id, inicio, fim))
    return cur.fetchall()


def _linhas_razao(cur, academia_id, inicio, fim):
    cur.execute(
        """SELECT COALESCE(l.data_efetivacao, l.data_competencia) AS data, l.id, c.nome AS conta,
                  l.descricao, l.origem, l.sentido, l.valor, fp.nome AS forma,
                  l.id_estorno, l.id_lancamento_origem
           FROM fin_lancamentos l
           JOIN fin_contas c ON c.id = l.id_conta
           LEFT JOIN formas_pagamento fp ON fp.id = l.id_forma_pagamento
           WHERE l.id_academia = %s AND l.status = 'efetivado'
             AND COALESCE(l.data_efetivacao, l.data_competencia) BETWEEN %s AND %s
           ORDER BY c.nome, data, l.id""", (academia_id, inicio, fim))
    return cur.fetchall()


ORIGENS = {"mensalidade": "Mensalidade", "cobranca_avulsa": "Cobrança avulsa", "receita": "Receita manual",
           "despesa": "Despesa", "estorno": "Estorno", "transferencia": "Transferência",
           "matricula": "Matrícula", "manual": "Manual"}


def planilha(cur, academia_id, academia_nome, inicio, fim):
    """Bytes do .xlsx."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    entradas = _linhas_entradas(cur, academia_id, inicio, fim)
    saidas = _linhas_saidas(cur, academia_id, inicio, fim)
    razao = _linhas_razao(cur, academia_id, inicio, fim)

    wb = Workbook()
    negrito = Font(bold=True)
    cab_fill = PatternFill("solid", fgColor="1C1C1C")
    cab_font = Font(bold=True, color="FFFFFF")

    def tabela(ws, cabecalho, linhas, moedas=(), datas=(), larguras=None):
        ws.append(cabecalho)
        for cel in ws[ws.max_row]:
            cel.fill, cel.font = cab_fill, cab_font
            cel.alignment = Alignment(vertical="center")
        ini = ws.max_row + 1
        for ln in linhas:
            ws.append(ln)
        for i in range(ini, ws.max_row + 1):
            for col in moedas:
                ws.cell(row=i, column=col).number_format = FMT_MOEDA
            for col in datas:
                ws.cell(row=i, column=col).number_format = FMT_DATA
        ws.freeze_panes = ws.cell(row=ini, column=1)
        if ws.max_row >= ini:
            ws.auto_filter.ref = f"A{ini - 1}:{get_column_letter(len(cabecalho))}{ws.max_row}"
        for i, w in enumerate(larguras or [], start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        return ini

    # ---------------- Entradas
    ws = wb.active
    ws.title = "Entradas"
    tabela(ws, ["Data", "Nº", "Descrição", "Categoria", "Aluno", "Forma", "Conta", "Valor"],
           [[r["data"], r["id"], r["descricao"], r["categoria"], r["aluno"] or "", r["forma"] or "Não informada",
             r["conta"] or "", float(r["valor"] or 0)] for r in entradas],
           moedas=(8,), datas=(1,), larguras=[12, 8, 46, 20, 32, 18, 16, 14])
    if entradas:
        ws.append(["Total", None, None, None, None, None, None, f"=SUM(H2:H{ws.max_row})"])
        ws.cell(row=ws.max_row, column=8).number_format = FMT_MOEDA
        for c in ws[ws.max_row]:
            c.font = negrito

    # ---------------- Saídas
    ws = wb.create_sheet("Saídas")
    tabela(ws, ["Data do pagamento", "Nº", "Vencimento", "Descrição", "Categoria", "Forma", "Conta", "Valor pago"],
           [[r["data"], r["id"], r["vencimento"], r["descricao"], r["categoria"], r["forma"] or "Não informada",
             r["conta"] or "", float(r["valor"] or 0)] for r in saidas],
           moedas=(8,), datas=(1, 3), larguras=[16, 8, 12, 46, 22, 18, 16, 14])
    if saidas:
        ws.append(["Total", None, None, None, None, None, None, f"=SUM(H2:H{ws.max_row})"])
        ws.cell(row=ws.max_row, column=8).number_format = FMT_MOEDA
        for c in ws[ws.max_row]:
            c.font = negrito

    # ---------------- Razão
    ws = wb.create_sheet("Razão")
    tabela(ws, ["Conta", "Data", "Nº", "Descrição", "Origem", "Forma", "Entrada", "Saída"],
           [[r["conta"], r["data"], r["id"], r["descricao"], ORIGENS.get(r["origem"], r["origem"] or ""),
             r["forma"] or "", float(r["valor"]) if r["sentido"] == "entrada" else None,
             float(r["valor"]) if r["sentido"] == "saida" else None] for r in razao],
           moedas=(7, 8), datas=(2,), larguras=[18, 12, 8, 46, 18, 16, 14, 14])

    # ---------------- Resumo (primeira aba)
    ws = wb.create_sheet("Resumo", 0)
    tot_e = sum((Decimal(str(r["valor"] or 0)) for r in entradas), Decimal("0"))
    tot_s = sum((Decimal(str(r["valor"] or 0)) for r in saidas), Decimal("0"))
    ws.append(["Movimento financeiro para a contabilidade"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([academia_nome])
    ws.append([f"Período: {inicio:%d/%m/%Y} a {fim:%d/%m/%Y}"])
    ws.append([])
    for rot, v in (("Entradas", tot_e), ("Saídas", tot_s), ("Resultado", tot_e - tot_s)):
        ws.append([rot, float(v)])
        ws.cell(row=ws.max_row, column=2).number_format = FMT_MOEDA
        ws.cell(row=ws.max_row, column=1).font = negrito

    def quebra(titulo, linhas, chave):
        soma = {}
        for r in linhas:
            k = r[chave] or "Não informada"
            soma[k] = soma.get(k, Decimal("0")) + Decimal(str(r["valor"] or 0))
        ws.append([])
        ws.append([titulo])
        ws.cell(row=ws.max_row, column=1).font = negrito
        for k, v in sorted(soma.items(), key=lambda kv: -kv[1]):
            ws.append([k, float(v)])
            ws.cell(row=ws.max_row, column=2).number_format = FMT_MOEDA

    quebra("Entradas por categoria", entradas, "categoria")
    quebra("Entradas por forma de pagamento", entradas, "forma")
    quebra("Saídas por categoria", saidas, "categoria")
    ws.append([])
    ws.append(["Critérios: entradas = receitas não canceladas, pela data; saídas = despesas pagas, "
               "pela data do pagamento; razão = lançamentos efetivados do caixa."])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), {"entradas": len(entradas), "saidas": len(saidas), "razao": len(razao)}
