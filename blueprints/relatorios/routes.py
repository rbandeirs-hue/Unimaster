# -*- coding: utf-8 -*-
"""Página de relatórios: um catálogo, três saídas.

O catálogo (`utils/relatorios.py`) diz o que cada relatório pergunta ao
banco. Aqui mora só o que é comum a todos: o recorte que vem da URL, a
tabela na tela, o Excel, o PDF e a folha de impressão. Acrescentar um
relatório não exige tocar neste arquivo.
"""

from datetime import date, timedelta
import calendar
import io

from flask import (Blueprint, render_template, request, redirect, url_for,
                   flash, abort, Response, session)
from flask_login import login_required, current_user

from config import get_db_connection
from utils.permissoes import somente_gestao
from utils.relatorios import (CATALOGO, Recorte, por_secao, contagem_faltas,
                              FALTAS, SECOES as SECOES_LISTA)

bp_relatorios = Blueprint("relatorios", __name__, url_prefix="/relatorios")
bp_relatorios.before_request(somente_gestao())


# ----------------------------------------------------------- recorte

def _academias_do_usuario():
    """Ids de academia que o usuário administra, na ordem em que aparecem."""
    ids = []
    try:
        from blueprints.academia.routes import _academias_ids_brutas
        ids = _academias_ids_brutas() or []
    except Exception:
        ids = []
    if not ids:
        atual = session.get("academia_gerenciamento_id") or getattr(current_user, "id_academia", None)
        ids = [atual] if atual else []
    return [int(i) for i in ids if i]


def _periodo():
    """Período da URL. `p` aceita atalhos; `de`/`ate` mandam quando vêm."""
    hoje = date.today()
    atalho = (request.args.get("p") or "mes").strip()
    de, ate = request.args.get("de"), request.args.get("ate")
    if de and ate:
        try:
            return (date.fromisoformat(de), date.fromisoformat(ate), "datas")
        except ValueError:
            pass
    if atalho == "mes_passado":
        fim = hoje.replace(day=1) - timedelta(days=1)
        return fim.replace(day=1), fim, atalho
    if atalho == "90d":
        return hoje - timedelta(days=89), hoje, atalho
    if atalho == "ano":
        return date(hoje.year, 1, 1), date(hoje.year, 12, 31), atalho
    ultimo = calendar.monthrange(hoje.year, hoje.month)[1]
    return hoje.replace(day=1), hoje.replace(day=ultimo), "mes"


def _academia_da_sessao(ids):
    """A academia desta sessão — a mesma que o resto do sistema usa.

    Não sai do filtro da tela: quem administra mais de uma já escolheu na
    entrada, e repetir a escolha aqui só criaria um caminho para pedir o
    número de uma academia que não é a sua.
    """
    atual = session.get("academia_gerenciamento_id") or getattr(current_user, "id_academia", None)
    if atual and (not ids or int(atual) in ids):
        return int(atual)
    return ids[0] if ids else None


def _recorte(cur):
    ids = _academias_do_usuario()
    academia_id = _academia_da_sessao(ids)
    if not academia_id:
        return None, ids
    ini, fim, atalho = _periodo()
    cur.execute("SELECT nome FROM academias WHERE id=%s", [academia_id])
    linha = cur.fetchone()
    rec = Recorte(
        academia_id=academia_id, ini=ini, fim=fim,
        modalidade_id=request.args.get("modalidade_id", type=int),
        turma_id=request.args.get("turma_id", type=int),
        academia_nome=(linha or {}).get("nome", ""),
    )
    rec.atalho = atalho
    rec.academias_ids = ids
    return rec, ids


def _filtros_da_url():
    """O recorte como querystring, para os links não perderem o filtro."""
    manter = ("p", "de", "ate", "modalidade_id", "turma_id")
    return {k: v for k, v in request.args.items() if k in manter and v}


# ----------------------------------------------------------- páginas

@bp_relatorios.route("/")
@login_required
def hub():
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        rec, ids = _recorte(cur)
        if not rec:
            flash("Nenhuma academia disponível para relatórios.", "warning")
            return redirect(url_for("painel.home"))

        faltas = contagem_faltas(cur, rec)
        cur.execute("""SELECT DISTINCT m.id, m.nome FROM modalidade m
                       JOIN academia_modalidades am ON am.modalidade_id = m.id
                       WHERE am.academia_id = %s ORDER BY m.nome""", [rec.academia_id])
        modalidades = cur.fetchall()
        cur.execute("SELECT TurmaID id, Nome nome FROM turmas WHERE id_academia=%s ORDER BY Nome",
                    [rec.academia_id])
        turmas = cur.fetchall()
    finally:
        cur.close()
        conn.close()

    grupos = por_secao()
    return render_template(
        "relatorios/hub.html",
        grupos=grupos, rec=rec, faltas=faltas, total=sum(len(g["itens"]) for g in grupos),
        modalidades=modalidades, turmas=turmas, filtros=_filtros_da_url(),
    )


def _executa(rid):
    """Roda um relatório e devolve (definicao, linhas, recorte)."""
    definicao = CATALOGO.get(rid)
    if not definicao:
        abort(404)
    if not definicao.pronto:
        return definicao, None, None
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        rec, _ = _recorte(cur)
        if not rec:
            return definicao, None, None
        linhas = definicao.funcao(cur, rec)
    finally:
        cur.close()
        conn.close()
    return definicao, linhas, rec


@bp_relatorios.route("/<rid>")
@login_required
def ver(rid):
    definicao, linhas, rec = _executa(rid)
    if linhas is None:
        flash(f"O relatório “{definicao.nome}” ainda está em construção.", "info")
        return redirect(url_for("relatorios.hub", **_filtros_da_url()))

    # O que a tela mostra em volta da tabela: trilha, fichas de filtro e, nos
    # relatórios de cadastro, o total da base para a frase "35 de 312".
    secao_titulo = next((t for c, t, _i, _r in SECOES_LISTA if c == definicao.secao), "Relatórios")
    conn = get_db_connection()
    cur = conn.cursor(dictionary=True)
    try:
        modalidade_nome = turma_nome = None
        if rec.modalidade_id:
            cur.execute("SELECT nome FROM modalidade WHERE id=%s", [rec.modalidade_id])
            r = cur.fetchone()
            modalidade_nome = r["nome"] if r else None
        if rec.turma_id:
            cur.execute("SELECT Nome nome FROM turmas WHERE TurmaID=%s", [rec.turma_id])
            r = cur.fetchone()
            turma_nome = r["nome"] if r else None
        base_total = None
        if definicao.secao == "cadastro":
            base_total = contagem_faltas(cur, rec)["total"]
    finally:
        cur.close()
        conn.close()

    return render_template(
        "relatorios/ver.html",
        d=definicao, linhas=linhas, rec=rec, filtros=_filtros_da_url(),
        totais=_totais(definicao, linhas), secao_titulo=secao_titulo,
        modalidade_nome=modalidade_nome, turma_nome=turma_nome, base_total=base_total,
        tem_ficha=bool(linhas and linhas[0].get("_aluno_id")),
        hoje_br=date.today().strftime("%d/%m/%Y"),
    )


def _totais(definicao, linhas):
    if not definicao.totaliza:
        return {}
    soma = {}
    for chave in definicao.totaliza:
        valores = [l.get(chave) for l in linhas]
        numeros = [v for v in valores if isinstance(v, (int, float))]
        soma[chave] = sum(numeros) if numeros else 0
    return soma


# ----------------------------------------------------------- saídas

def _cabecalho(definicao, rec):
    """As quatro linhas que todo arquivo leva: sem elas, papel solto na mesa
    não diz de que academia nem de que período é."""
    return [
        definicao.nome,
        f"{rec.academia_nome} · {rec.rotulo()}",
        f"Emitido em {date.today().strftime('%d/%m/%Y')} por {getattr(current_user, 'nome', '')}",
        definicao.criterio or "",
    ]


def _texto(valor, tipo):
    if valor is None or valor == "":
        return "—"
    if tipo == "moeda":
        return f"R$ {valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if tipo == "percentual":
        return f"{valor:.1f}%"
    if tipo == "data" and hasattr(valor, "strftime"):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


@bp_relatorios.route("/<rid>/excel")
@login_required
def excel(rid):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash("Biblioteca openpyxl não instalada — não dá para gerar o Excel.", "danger")
        return redirect(url_for("relatorios.ver", rid=rid, **_filtros_da_url()))

    definicao, linhas, rec = _executa(rid)
    if linhas is None:
        return redirect(url_for("relatorios.hub", **_filtros_da_url()))

    wb = Workbook()
    ws = wb.active
    ws.title = definicao.nome[:31]
    colunas = [c for c in definicao.colunas if not c.so_tela]

    for i, texto in enumerate(_cabecalho(definicao, rec), start=1):
        if not texto:
            continue
        ws.cell(row=i, column=1, value=texto).font = Font(bold=(i == 1), size=13 if i == 1 else 10,
                                                          color="1C1C1C" if i == 1 else "6B6862")
    inicio = 6
    for j, col in enumerate(colunas, start=1):
        c = ws.cell(row=inicio, column=j, value=col.rotulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C1C1C")
        c.alignment = Alignment(horizontal="left")

    for i, linha in enumerate(linhas, start=inicio + 1):
        for j, col in enumerate(colunas, start=1):
            valor = linha.get(col.chave)
            celula = ws.cell(row=i, column=j)
            if col.tipo in ("moeda", "numero", "percentual") and isinstance(valor, (int, float)):
                celula.value = valor
                celula.number_format = ('R$ #,##0.00' if col.tipo == "moeda"
                                        else '0.0"%"' if col.tipo == "percentual" else '0')
            elif col.tipo == "data" and hasattr(valor, "strftime"):
                celula.value = valor
                celula.number_format = "DD/MM/YYYY"
            else:
                celula.value = "" if valor is None else str(valor)

    # Coluna extra só na planilha: o link da ficha. Quem vai corrigir 35
    # cadastros não deve ter que procurar cada um pelo nome.
    if linhas and linhas[0].get("_aluno_id"):
        j = len(colunas) + 1
        c = ws.cell(row=inicio, column=j, value="Link da ficha")
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C1C1C")
        for i, linha in enumerate(linhas, start=inicio + 1):
            try:
                alvo = url_for("alunos.editar_aluno", aluno_id=linha["_aluno_id"], _external=True)
            except Exception:
                alvo = ""
            celula = ws.cell(row=i, column=j, value=alvo)
            if alvo:
                celula.hyperlink = alvo
                celula.style = "Hyperlink"
        ws.column_dimensions[get_column_letter(j)].width = 46

    for j, col in enumerate(colunas, start=1):
        largura = max([len(col.rotulo)] +
                      [len(_texto(l.get(col.chave), col.tipo)) for l in linhas[:200]] or [10])
        ws.column_dimensions[get_column_letter(j)].width = min(max(largura + 3, 12), 48)
    ws.freeze_panes = ws.cell(row=inicio + 1, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    nome = f"{rid}_{rec.ini:%Y%m%d}_{rec.fim:%Y%m%d}.xlsx"
    resp = Response(buf.getvalue(),
                    mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resp


@bp_relatorios.route("/<rid>/pdf")
@login_required
def pdf(rid):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    except ImportError:
        flash("Biblioteca reportlab não instalada. Use o Excel ou a impressão.", "warning")
        return redirect(url_for("relatorios.ver", rid=rid, **_filtros_da_url()))

    definicao, linhas, rec = _executa(rid)
    if linhas is None:
        return redirect(url_for("relatorios.hub", **_filtros_da_url()))

    colunas = [c for c in definicao.colunas if not c.so_tela]
    # Mais de seis colunas não cabe em retrato sem virar letra de bula.
    tamanho = landscape(A4) if len(colunas) > 6 else A4

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=tamanho, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm, title=definicao.nome)
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=15, alignment=0,
                            spaceAfter=2, textColor=colors.HexColor("#1C1C1C"))
    apoio = ParagraphStyle("s", parent=estilos["Normal"], fontSize=8.5,
                           textColor=colors.HexColor("#6B6862"), leading=12)
    celula = ParagraphStyle("c", parent=estilos["Normal"], fontSize=7.6, leading=9.6)

    cab = _cabecalho(definicao, rec)
    texto_cab = [Paragraph(cab[0], titulo), Paragraph(cab[1], apoio), Paragraph(cab[2], apoio)]
    if cab[3]:
        texto_cab.append(Paragraph(cab[3], apoio))

    # A logo da academia no canto: o relatório sai da sala e vira anexo de
    # e-mail, presta contas em reunião de associação. Sem a marca, é uma
    # tabela anônima que alguém precisa explicar de onde veio.
    logo = None
    try:
        from utils.contexto_logo import buscar_logo_arquivo
        caminho = buscar_logo_arquivo("academia", rec.academia_id)
        if caminho:
            from reportlab.platypus import Image as RLImage
            from PIL import Image as PILImage
            with PILImage.open(caminho) as img:
                larg, alt = img.size
            # Teto nos dois lados: uma logo deitada estourava a coluna e
            # saía cortada na borda da folha.
            altura, largura_max = 16 * mm, 30 * mm
            largura = altura * (larg / alt)
            if largura > largura_max:
                largura, altura = largura_max, largura_max * (alt / larg)
            logo = RLImage(caminho, width=largura, height=altura)
    except Exception:
        logo = None  # logo quebrada não pode impedir o relatório de sair

    if logo is not None:
        cabecalho = Table([[texto_cab, logo]],
                          colWidths=[tamanho[0] - 28 * mm - 34 * mm, 34 * mm])
        cabecalho.setStyle(TableStyle([
            ("VALIGN", (0, 0), (0, 0), "TOP"),
            ("VALIGN", (1, 0), (1, 0), "TOP"),
            ("ALIGN", (1, 0), (1, 0), "RIGHT"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ]))
        elementos = [cabecalho]
    else:
        elementos = list(texto_cab)
    elementos.append(Spacer(1, 7 * mm))

    dados = [[Paragraph(f"<b>{c.rotulo}</b>", celula) for c in colunas]]
    for l in linhas:
        dados.append([Paragraph(_texto(l.get(c.chave), c.tipo), celula) for c in colunas])

    if len(dados) == 1:
        elementos.append(Paragraph("Nada encontrado neste recorte.", apoio))
    else:
        largura = tamanho[0] - 28 * mm
        t = Table(dados, repeatRows=1, colWidths=[largura / len(colunas)] * len(colunas))
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1C1C1C")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D8D5D0")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAF9F7")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elementos.append(t)

    doc.build(elementos)
    buf.seek(0)
    nome = f"{rid}_{rec.ini:%Y%m%d}_{rec.fim:%Y%m%d}.pdf"
    resp = Response(buf.getvalue(), mimetype="application/pdf")
    resp.headers["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resp
